"""End-to-end and per-stage tests on the bundled 38-article sample dataset.

Link validation runs with HTTP checks disabled (see conftest.py), so results
are deterministic and need no network access.
"""
import json
import os

from openpyxl import load_workbook

from pipeline.config import DEFAULT_FMCG_CONFIG
from pipeline.dedup import dedup
from pipeline.ingestion import ingest
from pipeline.link_validator import validate_single
from pipeline.runner import run_pipeline

from .conftest import SAMPLE_CSV


def test_ingest_loads_and_normalises_sample():
    articles = ingest(SAMPLE_CSV, DEFAULT_FMCG_CONFIG)
    assert len(articles) == 38
    first = articles[0]
    for key in ("id", "title", "source", "published_date", "url", "full_text"):
        assert key in first


def test_exact_duplicates_are_collapsed():
    articles = ingest(SAMPLE_CSV, DEFAULT_FMCG_CONFIG)[:5]
    clone = dict(articles[0], id="dup-1")
    result = dedup(articles + [clone], DEFAULT_FMCG_CONFIG)
    titles = [a["title"] for a in result]
    assert titles.count(articles[0]["title"]) == 1


def test_link_validator_rejects_placeholder_and_malformed_urls():
    assert validate_single("https://example.com/story")["link_valid"] is False
    assert validate_single("not-a-url")["link_valid"] is False
    # Trusted domain + no network -> kept as "unverified" rather than dropped
    assert validate_single("https://www.reuters.com/markets/deals/")["link_valid"] is True


def test_full_pipeline_on_sample(tmp_path):
    result = run_pipeline(SAMPLE_CSV, str(tmp_path), DEFAULT_FMCG_CONFIG)
    s = result["summary"]

    # Funnel on the bundled sample (offline link validation)
    assert s["total_input"] == 38
    assert s["after_dedup"] == 30
    assert s["after_relevance"] == 24
    assert s["after_credibility"] == 23
    assert s["blocked"] == 1
    assert s["final_count"] == len(result["articles"]) == 20

    # Every surviving article is scored and ranked
    scores = [a["relevance_score"] for a in result["articles"]]
    assert scores == sorted(scores, reverse=True)
    assert all(a["link_valid"] for a in result["articles"])

    # All four output formats are written
    paths = result["output_paths"]
    assert set(paths) == {"json", "csv", "xlsx", "docx"}
    for p in paths.values():
        assert os.path.getsize(p) > 0

    with open(paths["json"], encoding="utf-8") as f:
        assert len(json.load(f)) == 20

    wb = load_workbook(paths["xlsx"])
    assert len(wb.sheetnames) >= 4

    assert result["newsletter"]
